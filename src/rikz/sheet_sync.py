"""Keep the hand-built portfolio workbook (an .xlsx on Google Drive) in step
with the website: after every new report the Awaed and Manafa input rows are
updated, and everything else in the workbook (formulas, charts, comments,
formatting, other tabs) is left exactly as it is.

The workbook belongs to its owner, so this is an upsert, never a rewrite:

* Manafa: a row is found by its Opportunity ID. An existing row is touched
  only when its status changes (Active -> Closed and so on); a closed row then
  gets its payment date and the gross profit, fee and VAT actually paid. Mandate
  positions that are not in the sheet yet are appended.
* Awaed: the sheet lists the deposits in order with no order ID, so rows are
  lined up with the confirmations by order date; the line-up is checked
  against principal and profit before anything is written. Existing rows only
  get their status updated; new confirmations are appended.

Cells are written by editing the sheet XML in place (openpyxl would drop the
workbook's charts on save), and the workbook is told to recalculate on open.
"""

from __future__ import annotations

import io
import logging
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from xml.sax.saxutils import escape

log = logging.getLogger("rikz.sheet_sync")
FIRST_ROW, LAST_ROW = 6, 65  # input rows; the formula columns are filled down to LAST_ROW
EXCEL_EPOCH = date(1899, 12, 30)
STATUS = {"active": "Active", "delayed": "Delay", "defaulted": "Default",
          "repaid": "Closed", "repaid_early": "Closed", "closed": "Closed"}


@dataclass
class Plan:
    edits: dict[str, dict[str, object]] = field(default_factory=dict)  # sheet -> {"B27": value}
    changes: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)

    def set(self, sheet: str, ref: str, value) -> None:
        self.edits.setdefault(sheet, {})[ref] = value


# -- planning -----------------------------------------------------------------
def _num(v) -> float | None:
    return float(v) if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool) else None


def _day(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


def plan(workbook: bytes, batch, rule) -> Plan:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(workbook))  # read only; never saved
    out = Plan()
    if batch.matching is not None:
        _plan_manafa(wb["Manafa"], batch, rule, out)
    if batch.chain is not None:
        _plan_awaed(wb["Awaed"], batch, out)
    return out


def _rows(ws, key_col: str) -> dict[int, object]:
    return {r: ws[f"{key_col}{r}"].value for r in range(FIRST_ROW, LAST_ROW + 1)
            if ws[f"A{r}"].value not in (None, "")}


def _plan_manafa(ws, batch, rule, out: Plan) -> None:
    from .manafa.mandate import in_mandate

    by_oid: dict[str, list] = {}
    for link in batch.matching.links:
        if in_mandate(link.position, rule):
            by_oid.setdefault(link.position.oid, []).append(link)
    existing = {v: r for r, v in _rows(ws, "B").items() if isinstance(v, str)}
    used = list(_rows(ws, "A"))
    next_row = (max(used) + 1) if used else FIRST_ROW
    next_no = int(max((_num(ws[f"A{r}"].value) or 0) for r in used) + 1) if used else 1

    for oid, links in sorted(by_oid.items(), key=lambda kv: min(l.position.entry for l in kv[1])):
        ps = [l.position for l in links]
        # a split investment is one row in the sheet: add its parts together
        status = "Closed" if all(STATUS[p.status.value] == "Closed" for p in ps) else \
            next(STATUS[p.status.value] for p in ps if STATUS[p.status.value] != "Closed")
        principal = sum((p.principal for p in ps), Decimal("0"))
        paid = max((l.paid_on for l in links if l.paid_on), default=None)
        if status == "Closed":
            gross = sum((l.settlement.gross_profit for l in links if l.settlement), Decimal("0"))
            fee = -sum((l.settlement.fee for l in links if l.settlement), Decimal("0"))
            vat = -sum((l.settlement.vat for l in links if l.settlement), Decimal("0"))
        else:  # expected figures, the way the sheet records open positions (20% fee, 15% VAT on it)
            net = sum((p.net_profit for p in ps), Decimal("0"))
            gross = (net / Decimal("0.77")).quantize(Decimal("0.01"))
            fee = (gross * Decimal("0.2")).quantize(Decimal("0.01"))
            vat = (fee * Decimal("0.15")).quantize(Decimal("0.01"))

        row = existing.get(oid)
        if row is not None:
            was = ws[f"I{row}"].value
            if was == status:
                continue
            out.set("Manafa", f"I{row}", status)
            text = f"Manafa {oid}: {was} → {status}"
            if status == "Closed":
                out.set("Manafa", f"J{row}", paid)
                out.set("Manafa", f"L{row}", float(gross))
                out.set("Manafa", f"M{row}", float(fee))
                out.set("Manafa", f"N{row}", float(vat))
                text += f", paid {paid}, gross profit {gross:,.2f}"
            out.changes.append(text)
            continue
        if next_row > LAST_ROW:
            out.skipped.append(f"Manafa {oid}: the sheet's formulas end at row {LAST_ROW}; extend them to add more")
            continue
        p0 = min(ps, key=lambda p: p.entry)
        r = next_row
        for col, value in (("A", next_no), ("B", oid), ("C", "Reinvested Capital"), ("D", float(p0.rate)),
                           ("E", p0.tenor_months * 30), ("F", p0.entry), ("G", p0.maturity), ("H", p0.rating),
                           ("I", status), ("J", paid), ("K", float(principal)), ("L", float(gross)),
                           ("M", float(fee)), ("N", float(vat))):
            out.set("Manafa", f"{col}{r}", value)
        out.changes.append(f"Manafa {oid}: added in row {r} ({principal:,.2f}, {status})")
        next_row += 1
        next_no += 1


def _plan_awaed(ws, batch, out: Plan) -> None:
    deposits = sorted(batch.chain.deposits, key=lambda d: d.ordered_at)
    as_of = batch.as_of
    rows = sorted(_rows(ws, "A"))
    if len(rows) > len(deposits):
        out.skipped.append(f"Awaed: the sheet has {len(rows)} rows but only {len(deposits)} confirmations are "
                           "stored; Awaed left unchanged")
        return
    for r, d in zip(rows, deposits):
        j, k = _num(ws[f"J{r}"].value), _num(ws[f"K{r}"].value)
        if (j is not None and abs(j - float(d.principal)) > 0.5) or \
                (k is not None and abs(k - float(d.total_return)) > 0.05):
            out.skipped.append(f"Awaed: row {r} does not line up with confirmation {d.order_id} "
                               f"({d.principal:,.2f}); Awaed left unchanged")
            return
    for r, d in zip(rows, deposits):
        status = "Active" if d.maturity > as_of else "Closed"
        was = ws[f"I{r}"].value
        if was != status:
            out.set("Awaed", f"I{r}", status)
            out.changes.append(f"Awaed row {r} ({d.principal:,.2f}): {was} → {status}")
    next_row = (rows[-1] + 1) if rows else FIRST_ROW
    next_id = int(max((_num(ws[f"A{r}"].value) or 0) for r in rows) + 1) if rows else 1
    for d in deposits[len(rows):]:
        if next_row > LAST_ROW:
            out.skipped.append(f"Awaed {d.order_id}: the sheet's formulas end at row {LAST_ROW}")
            continue
        rate = (d.total_return * 360 / (d.principal * d.tenor_days)).quantize(Decimal("0.0001"))
        status = "Active" if d.maturity > as_of else "Closed"
        for col, value in (("A", next_id), ("B", float(rate)), ("C", d.order_date), ("D", d.order_date),
                           ("E", d.maturity), ("F", d.tenor_days), ("G", "A+"), ("H", "EOP"), ("I", status),
                           ("J", float(d.principal)), ("K", float(d.total_return.quantize(Decimal("0.01")))),
                           ("L", float(d.fees)), ("M", float(d.vat))):
            out.set("Awaed", f"{col}{next_row}", value)
        out.changes.append(f"Awaed: added confirmation {d.order_id} in row {next_row} ({d.principal:,.2f}, {status})")
        next_row += 1
        next_id += 1


# -- writing the xlsx -----------------------------------------------------------
def _col_index(ref: str) -> int:
    n = 0
    for ch in re.match(r"[A-Z]+", ref).group():
        n = n * 26 + ord(ch) - 64
    return n


def _cell_xml(ref: str, style: str | None, value) -> str:
    s = f' s="{style}"' if style else ""
    if value is None:
        return f'<c r="{ref}"{s}/>'
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return f'<c r="{ref}"{s}><v>{(value - EXCEL_EPOCH).days}</v></c>'
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return f'<c r="{ref}"{s}><v>{value}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t>{escape(str(value))}</t></is></c>'


_CELL = r'<c r="{ref}"(?P<attrs>[^>]*?)(?:/>|>.*?</c>)'
_DATE_FMT_IDS = set(range(14, 23)) | set(range(45, 48))


def date_styles(styles_xml: str) -> set[str]:
    """Indexes of the cell styles that display a date."""
    custom = {m.group(1) for m in re.finditer(r'<numFmt [^>]*?numFmtId="(\d+)"[^>]*?formatCode="([^"]*)"', styles_xml)
              if re.search(r"[dmy]", re.sub(r'"[^"]*"|\[[^\]]*\]', "", m.group(2)), flags=re.I)}
    xfs = re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles_xml, flags=re.S)
    out = set()
    for i, m in enumerate(re.finditer(r"<xf\b[^>]*>", xfs.group(1) if xfs else "")):
        fid = re.search(r'numFmtId="(\d+)"', m.group(0))
        if fid and (int(fid.group(1)) in _DATE_FMT_IDS or fid.group(1) in custom):
            out.add(str(i))
    return out


def _date_style_near(xml: str, rownum: int, dates: set[str]) -> str | None:
    row = re.search(rf'<row r="{rownum}"[^>]*>(.*?)</row>', xml, flags=re.S)
    for scope in ((row.group(1),) if row else ()) + (xml,):
        for m in re.finditer(r'<c r="[A-Z]+\d+"[^>]*?\bs="(\d+)"', scope):
            if m.group(1) in dates:
                return m.group(1)
    return None


def patch_sheet(xml: str, cells: dict[str, object], dates: set[str] = frozenset()) -> str:
    for ref, value in sorted(cells.items(), key=lambda kv: (int(re.sub(r"[A-Z]", "", kv[0])), _col_index(kv[0]))):
        rownum = int(re.sub(r"[A-Z]", "", ref))
        is_date = isinstance(value, date)
        m = re.search(_CELL.format(ref=ref), xml, flags=re.S)
        if m:
            style = re.search(r'\bs="(\d+)"', m.group("attrs"))
            style = style.group(1) if style else None
            if is_date and style not in dates:
                style = _date_style_near(xml, rownum, dates) or style
            xml = xml[:m.start()] + _cell_xml(ref, style, value) + xml[m.end():]
            continue
        row = re.search(rf'<row r="{rownum}"(?P<attrs>[^>]*?)(?P<close>/>|>(?P<body>.*?)</row>)', xml, flags=re.S)
        new = _cell_xml(ref, _date_style_near(xml, rownum, dates) if is_date else None, value)
        if row is None:  # insert a new row before the first row with a higher number
            later = next((m for m in re.finditer(r'<row r="(\d+)"', xml) if int(m.group(1)) > rownum), None)
            at = later.start() if later else xml.index("</sheetData>")
            xml = xml[:at] + f'<row r="{rownum}">{new}</row>' + xml[at:]
            continue
        if row.group("close") == "/>":
            xml = xml[:row.start()] + f'<row r="{rownum}"{row.group("attrs")}>{new}</row>' + xml[row.end():]
            continue
        body_start = row.start("body")
        after = next((c for c in re.finditer(r'<c r="([A-Z]+)\d+"', row.group("body"))
                      if _col_index(c.group(1)) > _col_index(ref)), None)
        at = body_start + after.start() if after else row.end("body")
        xml = xml[:at] + new + xml[at:]
    return xml


def apply(workbook: bytes, edits: dict[str, dict[str, object]]) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(workbook))
    wbxml = zin.read("xl/workbook.xml").decode()
    rels = zin.read("xl/_rels/workbook.xml.rels").decode()
    targets = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship [^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', rels)}
    targets.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship [^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
    paths = {}
    for m in re.finditer(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="([^"]+)"', wbxml):
        t = targets[m.group(2)].lstrip("/")
        paths[m.group(1)] = t if t.startswith("xl/") else "xl/" + t
    dates = date_styles(zin.read("xl/styles.xml").decode()) if "xl/styles.xml" in zin.namelist() else set()
    changed = {paths[name]: patch_sheet(zin.read(paths[name]).decode(), cells, dates)
               for name, cells in edits.items() if cells}
    # recalculate every formula when the workbook is next opened
    if "<calcPr" in wbxml:
        wbxml = re.sub(r"<calcPr([^>]*?)/?>", lambda m: "<calcPr" + re.sub(r'\s*fullCalcOnLoad="[^"]*"', "", m.group(1))
                       + ' fullCalcOnLoad="1"/>', wbxml, count=1)
    else:
        wbxml = wbxml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
    changed["xl/workbook.xml"] = wbxml
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = changed[item.filename].encode() if item.filename in changed else zin.read(item.filename)
            zout.writestr(item, data)
    return buf.getvalue()


# -- running it -----------------------------------------------------------------
def configured() -> bool:
    return bool(os.environ.get("SHEET_SYNC_FILE_ID") and os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON"))


def sync(store, drive=None, file_id: str | None = None) -> dict:
    """Bring the workbook up to date with the latest report. Returns a summary
    (also kept in the 'sheet_sync_last' meta value for the admin page)."""
    import json

    from .ingest.drive import GoogleDrive
    from .store.db import now, set_meta

    file_id = file_id or os.environ["SHEET_SYNC_FILE_ID"]
    drive = drive or GoogleDrive.from_env(write=True)
    batch, rule = store.latest_batch()
    if batch is None:
        summary = {"at": now().isoformat(), "ok": True, "changes": [], "skipped": ["no report yet"]}
    else:
        data = drive.download_id(file_id)
        p = plan(data, batch, rule)
        if p.edits:
            drive.upload_id(file_id, apply(data, p.edits))
        summary = {"at": now().isoformat(), "ok": True, "as_of": str(batch.as_of),
                   "changes": p.changes, "skipped": p.skipped}
    set_meta(store.Session, "sheet_sync_last", json.dumps(summary))
    log.info("sheet sync: %s", summary)
    return summary


def attach(store) -> None:
    """Sync after every new report; if Drive is unreachable, a background job
    retries it."""
    from .ingest.worker import enqueue, handler

    @handler("sheet.sync")
    def _job(store_, payload, context):
        sync(store_)

    def on_result(result, source):
        if not result.snapshot_version:
            return
        try:
            sync(store)
        except Exception as exc:  # noqa: BLE001 - the report itself is already safe
            import json

            from .store.db import now, set_meta

            log.warning("sheet sync failed, queued for retry: %s", exc)
            set_meta(store.Session, "sheet_sync_last", json.dumps(
                {"at": now().isoformat(), "ok": False, "error": f"{type(exc).__name__}: {exc}"[:500]}))
            enqueue(store.Session, "sheet.sync", {})

    store.listeners.append(on_result)
