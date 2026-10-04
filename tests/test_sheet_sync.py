import io
import json
import zipfile

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from conftest import AWAED, MANAFA
from rikz import sheet_sync
from rikz.store.db import get_meta, init_db, make_engine
from rikz.store.files import FileStore
from rikz.store.service import Store


def read(paths):
    return [(p.name, p.read_bytes()) for p in paths]


@pytest.fixture(scope="module")
def store(tmp_path_factory):
    t = tmp_path_factory.mktemp("sync")
    st = Store(init_db(make_engine(f"sqlite:///{t / 'db.sqlite'}")), FileStore(t / "files"))
    st.ingest(read(sorted(AWAED.glob("*.pdf")) + [MANAFA / "2026-09-24_portfolio.xlsx",
                                                     MANAFA / "2026-09-24_account_statement.xlsx"]), "cli")
    return st


def workbook(manafa_rows, awaed_rows) -> bytes:
    """A small stand-in for the owner's workbook: input rows, formulas filled
    down to row 65, a summary tab and a chart."""
    wb = openpyxl.Workbook()
    wb.active.title = "Dashboard"
    ws = wb.create_sheet("Awaed")
    for r in range(6, 66):
        ws[f"N{r}"] = f'=IF($A{r}="","",K{r}-L{r}-M{r})'
    for i, row in enumerate(awaed_rows):
        for c, v in zip("ABCDEFGHIJKLM", row):
            ws[f"{c}{6 + i}"] = v
    ws = wb.create_sheet("Manafa")
    for r in range(6, 66):
        ws[f"O{r}"] = f'=IF($A{r}="","",L{r}-M{r}-N{r})'
    for i, row in enumerate(manafa_rows):
        for c, v in zip("ABCDEFGHIJKLMN", row):
            ws[f"{c}{6 + i}"] = v
    dash = wb["Dashboard"]
    dash["A1"], dash["A2"] = "Manafa net", "=SUM(Manafa!O6:O65)"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=11, min_row=6, max_row=30))
    dash.add_chart(chart, "C3")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FakeDrive:
    def __init__(self, data):
        self.data, self.uploads = data, 0

    def download_id(self, file_id):
        return self.data

    def upload_id(self, file_id, data):
        self.data, self.uploads = data, self.uploads + 1


def test_upsert_updates_changed_rows_appends_new_ones_and_keeps_the_rest(store):
    batch, rule = store.latest_batch()
    first = sorted(batch.chain.deposits, key=lambda d: d.ordered_at)[:2]
    awaed = [[i + 1, 0.04, d.order_date, d.order_date, d.maturity, d.tenor_days, "A+", "EOP", "Active",
              float(d.principal), round(float(d.total_return), 2), 0, 0] for i, d in enumerate(first)]
    closed = next(l for l in batch.matching.links
                  if l.settlement and l.position.principal >= rule.min_principal and l.position.entry >= rule.funded_from)
    p = closed.position
    manafa = [[1, p.oid, "Capital Injection", float(p.rate), 90, p.entry, p.maturity, p.rating, "Active", None,
               float(p.principal), 1.0, 0.2, 0.03]]
    drive = FakeDrive(workbook(manafa, awaed))
    before = drive.data

    summary = sheet_sync.sync(store, drive=drive, file_id="sheet-id")
    assert summary["ok"] and drive.uploads == 1
    assert any(f"Manafa {p.oid}: Active → Closed" in c for c in summary["changes"])
    assert any("Awaed row 6" in c and "→ Closed" in c for c in summary["changes"])
    assert any(c.startswith("Awaed: added confirmation") for c in summary["changes"])

    wb = openpyxl.load_workbook(io.BytesIO(drive.data))
    m = wb["Manafa"]
    assert m["I6"].value == "Closed" and m["J6"].value.date() == closed.paid_on
    assert m["L6"].value == float(closed.settlement.gross_profit)
    assert m["C6"].value == "Capital Injection"  # the owner's own entry is kept
    assert m["B7"].value and m["B7"].value != p.oid  # other mandate positions appended
    assert m["O7"].value == '=IF($A7="","",L7-M7-N7)'  # formulas untouched
    a = wb["Awaed"]
    n = len(batch.chain.deposits)
    assert a[f"A{5 + n}"].value == n and a[f"G{5 + n}"].value == "A+"

    # only the two input tabs and the recalculation flag changed; the chart is still there
    zin, zout = zipfile.ZipFile(io.BytesIO(before)), zipfile.ZipFile(io.BytesIO(drive.data))
    assert zin.namelist() == zout.namelist()
    changed = {x for x in zin.namelist() if zin.read(x) != zout.read(x)}
    assert changed == {"xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml", "xl/workbook.xml"}
    assert any(x.startswith("xl/charts/") for x in zout.namelist())
    assert b'fullCalcOnLoad="1"' in zout.read("xl/workbook.xml")
    assert json.loads(get_meta(store.Session, "sheet_sync_last"))["changes"] == summary["changes"]

    # a second run has nothing left to do and does not upload
    again = sheet_sync.sync(store, drive=drive, file_id="sheet-id")
    assert again["changes"] == [] and drive.uploads == 1


def test_awaed_rows_that_do_not_line_up_are_left_alone(store):
    wrong = [[1, 0.04, None, None, None, 7, "A+", "EOP", "Active", 999.0, 1.0, 0, 0]]
    drive = FakeDrive(workbook([], wrong))
    summary = sheet_sync.sync(store, drive=drive, file_id="sheet-id")
    assert any("does not line up" in s for s in summary["skipped"])
    assert not any(c.startswith("Awaed") for c in summary["changes"])
    assert openpyxl.load_workbook(io.BytesIO(drive.data))["Awaed"]["J6"].value == 999.0


def test_sync_runs_after_each_new_report_and_failures_are_retried(tmp_path, monkeypatch):
    from rikz.store.db import Job

    st = Store(init_db(make_engine(f"sqlite:///{tmp_path / 'db.sqlite'}")), FileStore(tmp_path / "files"))
    calls = []
    monkeypatch.setattr(sheet_sync, "sync", lambda store, **kw: calls.append(1) or (_ for _ in ()).throw(OSError("down")))
    sheet_sync.attach(st)
    st.ingest(read(sorted(AWAED.glob("*.pdf")) + [MANAFA / "2026-09-24_portfolio.xlsx",
                                                     MANAFA / "2026-09-24_account_statement.xlsx"]), "upload")
    assert calls == [1]
    assert not json.loads(get_meta(st.Session, "sheet_sync_last"))["ok"]
    with st.Session() as s:
        assert [j.kind for j in s.query(Job).all()] == ["sheet.sync"]
